from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.reports.services.pdf_cabecera import cabecera_flujo

AZUL = '1F3864'
ROJO = 'C8102E'
ANCHO_UTILIZABLE = A4[0] - 30 * mm


# =============================================================================
# PDF
# =============================================================================
def generar_pdf_auditoria(header, tablas):
    """Documento PDF con la cabecera de auditoría y sus tablas de detalle."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        rightMargin=15 * mm, leftMargin=15 * mm,
        topMargin=15 * mm, bottomMargin=18 * mm,
        title=header['titulo'], author=header['generado_por'],
    )

    estilos = _estilos_pdf()
    historia = []

    historia += cabecera_flujo(header, estilos)
    historia.append(_tabla_metadatos_pdf(header, estilos))
    historia.append(Spacer(1, 6 * mm))

    for entrada in tablas:
        if isinstance(entrada, dict):
            if entrada.get('solo_excel'):
                continue
            historia += _seccion_tarjetas(entrada, estilos)
            continue
        titulo, columnas, filas = entrada
        historia.append(Paragraph(titulo, estilos['subtitulo']))
        historia.append(Spacer(1, 2 * mm))
        historia.append(_tabla_generica_pdf(columnas, filas, estilos['celda']))
        historia.append(Spacer(1, 5 * mm))

    doc.build(historia, onFirstPage=_pie_pagina, onLaterPages=_pie_pagina)
    buffer.seek(0)
    return buffer


def _estilos_pdf():
    base = getSampleStyleSheet()
    return {
        'marca': ParagraphStyle('marca', parent=base['Normal'],
                                fontName='Helvetica-Bold', fontSize=10,
                                textColor=colors.HexColor('#' + ROJO)),
        'titulo': ParagraphStyle('titulo', parent=base['Title'], fontSize=16,
                                 leading=20),
        'subtitulo': ParagraphStyle('subtitulo', parent=base['Heading2'],
                                    fontSize=12, spaceAfter=6,
                                    textColor=colors.HexColor('#' + ROJO)),
        'celda': ParagraphStyle('celda', parent=base['Normal'], fontSize=8.5,
                                leading=11),
    }


def _tabla_metadatos_pdf(header, estilos):
    filas = []
    for etiqueta, valor in (
        ('Rango de fechas', header['rango']),
        ('Sede(s)', header['sede']),
        ('Registros', header.get('registros', 0)),
        ('Generado por', header['generado_por']),
        ('Generado el', header['generado_en']),
    ):
        filas.append([etiqueta, str(valor)])

    tabla = Table(filas, colWidths=[40 * mm, ANCHO_UTILIZABLE - 40 * mm])
    tabla.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#8A0F1D')),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#FBE9EA')),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#E8CDCF')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    return tabla


def _tabla_generica_pdf(columnas, filas, estilo_celda):
    if not filas:
        filas = [['Sin registros para el rango seleccionado']]
        columnas = ['']
    celdas = [[Paragraph(str(c), estilo_celda) for c in columnas]]
    for fila in filas:
        celdas.append([Paragraph(str(v), estilo_celda) for v in fila])

    ancho_celda = ANCHO_UTILIZABLE / max(len(columnas), 1)
    tabla = Table(celdas, colWidths=[ancho_celda] * len(columnas))
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#' + ROJO)),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#E8CDCF')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1),
         [colors.white, colors.HexColor('#FFF6F6')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    return tabla


def _pie_pagina(canvas, doc):
    from datetime import datetime
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor('#' + ROJO))
    canvas.setLineWidth(0.9)
    canvas.line(15 * mm, 13.5 * mm, A4[0] - 15 * mm, 13.5 * mm)
    canvas.setFont('Helvetica', 7.5)
    canvas.setFillColor(colors.grey)
    canvas.drawString(15 * mm, 10 * mm,
                      f"Sistema PH — Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, f'Página {doc.page}')
    canvas.restoreState()


# =============================================================================
# INFORME TIPO TARJETA (fichas que replican los listados web)
# =============================================================================
def _seccion_tarjetas(entrada, estilos):
    """Sección de tarjetas: una ficha por registro (evento, compra, traslado...).
    """
    flujos = [Paragraph(entrada['titulo'], estilos['subtitulo']),
              Spacer(1, 2 * mm)]
    if not entrada.get('items'):
        flujos.append(Paragraph('Sin registros para el rango seleccionado.',
                                estilos['celda']))
        flujos.append(Spacer(1, 5 * mm))
        return flujos
    for item in entrada['items']:
        flujos.append(_tarjeta(item))
        flujos.append(Spacer(1, 3 * mm))
    return flujos


def _estilos_tarjeta():
    base = getSampleStyleSheet()
    return {
        'cabeza': ParagraphStyle('cabeza', parent=base['Normal'],
                                 fontName='Helvetica-Bold', fontSize=9,
                                 leading=12, textColor=colors.white),
        'cabeza_derecha': ParagraphStyle('cabeza_derecha',
                                         parent=base['Normal'], fontSize=7.5,
                                         leading=12, textColor=colors.white,
                                         alignment=TA_RIGHT),
        'etiqueta': ParagraphStyle('etiqueta', parent=base['Normal'],
                                   fontName='Helvetica-Bold', fontSize=8,
                                   leading=10,
                                   textColor=colors.HexColor('#8A0F1D')),
        'valor': ParagraphStyle('valor', parent=base['Normal'], fontSize=8,
                                leading=11,
                                textColor=colors.HexColor('#1F2328')),
    }


def _tarjeta(item):
    """Tarjeta individual: cabecera roja (acción + fecha) y pares etiqueta/
    valor. Recibe un dict con 'banda', opcional 'derecha' y 'filas' como
    lista de tuplas (etiqueta, valor)."""
    estilos = _estilos_tarjeta()
    ancho_barra = 2.5 * mm
    ancho_etiqueta = 38 * mm
    ancho_valor = ANCHO_UTILIZABLE - ancho_barra - ancho_etiqueta

    banda = escape(item.get('banda') or '—')
    derecha = escape(item.get('derecha') or '')

    datos = [['', Paragraph(banda, estilos['cabeza']),
              Paragraph(derecha, estilos['cabeza_derecha']) if derecha else '']]
    for etiqueta, valor in item.get('filas') or []:
        datos.append(['', Paragraph(escape(etiqueta), estilos['etiqueta']),
                      Paragraph(escape(str(valor)), estilos['valor'])])

    tabla = Table(datos, colWidths=[ancho_barra, ancho_etiqueta, ancho_valor])
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (2, 0), colors.HexColor('#' + ROJO)),
        ('BACKGROUND', (0, 1), (0, -1), colors.HexColor('#FBE9EA')),
        ('BOX', (0, 0), (-1, -1), 0.6, colors.HexColor('#E8CDCF')),
        ('INNERGRID', (0, 1), (-1, -1), 0.3, colors.HexColor('#F2E2E3')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (1, 0), (2, 0), 8),
        ('RIGHTPADDING', (0, 0), (0, 0), 2 * mm),
    ]))
    return tabla


# =============================================================================
# EXCEL
# =============================================================================
def generar_excel_auditoria(header, tablas):
    """Libro Excel con la cabecera de auditoría y sus tablas de detalle."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Auditoria'

    hoja['A1'] = 'Sistema PH'
    hoja['A1'].font = Font(bold=True, size=14, color=ROJO)
    hoja['A2'] = header['titulo']
    hoja['A2'].font = Font(bold=True, size=11, color=AZUL)
    _fila_vacia(hoja)

    renglones = [
        ('Rango de fechas', header['rango']),
        ('Sede(s)', header['sede']),
        ('Registros', header.get('registros', 0)),
        ('Generado por', header['generado_por']),
        ('Generado el', header['generado_en']),
    ]
    for etiqueta, valor in renglones:
        hoja.append([etiqueta, valor])
        hoja[hoja.max_row][0].font = Font(bold=True, size=9)
        hoja[hoja.max_row][0].alignment = Alignment(vertical='top')
        hoja[hoja.max_row][1].font = Font(size=9)
        hoja[hoja.max_row][1].alignment = Alignment(vertical='top',
                                                     wrap_text=True)
    _fila_vacia(hoja)

    for entrada in tablas:
        if isinstance(entrada, dict):
            fila_titulo = hoja.max_row + 1
            hoja.append([entrada['titulo']])
            n_columnas = _tabla_datos_excel(hoja, entrada['columnas'],
                                            entrada['filas'])
        else:
            titulo, columnas, filas = entrada
            fila_titulo = hoja.max_row + 1
            hoja.append([titulo])
            celda = hoja[fila_titulo][0]
            celda.font = Font(bold=True, size=10, color=AZUL)
            n_columnas = _tabla_datos_excel(hoja, columnas, filas)
        if n_columnas > 1:
            hoja.merge_cells(
                start_row=fila_titulo, start_column=1,
                end_row=fila_titulo, end_column=n_columnas)
        _fila_vacia(hoja)

    hoja.append(['Documento generado por Sistema PH'])
    hoja.append([f"Generado el {header['generado_en']} por {header['generado_por']}"])
    for celda in hoja[hoja.max_row - 1]:
        celda.font = Font(size=8, italic=True, color='808080')
    for celda in hoja[hoja.max_row]:
        celda.font = Font(size=8, italic=True, color='808080')

    _ajustar_columnas(hoja)

    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer


def _fila_vacia(hoja):
    hoja.append([])


def _tabla_datos_excel(hoja, columnas, filas):
    encabezados = columnas if columnas else ['Detalle']
    hoja.append(encabezados)
    fila_enc = hoja.max_row
    for celda in hoja[fila_enc]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor=AZUL)
        celda.alignment = Alignment(horizontal='center', vertical='center',
                                    wrap_text=True)

    filas_mostrar = filas if filas else [['Sin registros para el rango seleccionado']]
    for fila in filas_mostrar:
        hoja.append(fila)
        fila_actual = hoja.max_row
        for celda in hoja[fila_actual]:
            celda.font = Font(size=9)
            celda.alignment = Alignment(vertical='top', wrap_text=True)

    return len(encabezados)


def _ajustar_columnas(hoja, minimo=12, maximo=48, margen=3):
    """Ajusta el ancho de cada columna al contenido más largo.

    El texto largo queda envuelto dentro de la celda (wrap_text) en vez de
    desbordarse, y el alto de fila lo recalcula Excel al abrir el archivo.
    """
    longitudes = {}
    for fila in hoja.iter_rows():
        for celda in fila:
            if celda.value is None:
                continue
            lineas = str(celda.value).split('\n')
            largo = max((len(linea) for linea in lineas), default=0)
            estimado = largo * 1.1 + margen
            col = celda.column - 1
            if estimado > longitudes.get(col, 0):
                longitudes[col] = estimado
    for col, largo in longitudes.items():
        hoja.column_dimensions[get_column_letter(col + 1)].width = (
            min(max(largo, minimo), maximo))